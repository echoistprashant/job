import io
import pytest
import pymupdf
import docx
from fastapi.testclient import TestClient
from backend.app.main import app

client = TestClient(app)


def create_sample_pdf_bytes() -> bytes:
    """Create a minimal valid PDF with realistic resume text in-memory."""
    doc = pymupdf.open()
    page = doc.new_page()
    text = (
        "Prashant Yadav\n"
        "prashant@example.com | +91 9876543210\n"
        "https://github.com/echoistprashant | https://linkedin.com/in/prashantyadav\n\n"
        "SUMMARY\n"
        "Passionate AI & Backend Engineer with experience building scalable systems and agentic workflows.\n\n"
        "EDUCATION\n"
        "Bachelor of Technology in Computer Science and Engineering\n\n"
        "SKILLS\n"
        "Python, FastAPI, SQLAlchemy, PostgreSQL, Docker, LangChain, LangGraph, LLM, RAG, PyTorch\n\n"
        "EXPERIENCE\n"
        "AI Engineer at TechCorp\n"
        "Built intelligent job aggregation and matching pipelines using LLMs and vector embeddings.\n\n"
        "PROJECTS\n"
        "Project: AI Job Application Agent\n"
        "Designed autonomous application system with human approval.\n"
    )
    page.insert_text((50, 72), text, fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes() -> bytes:
    """Create a minimal valid DOCX with realistic resume text in-memory."""
    doc = docx.Document()
    doc.add_heading("Alex Johnson", level=1)
    doc.add_paragraph("alex.johnson@example.com | +1 234-567-8900")
    doc.add_paragraph("Summary: Experienced backend software engineer passionate about microservices.")
    doc.add_heading("Education", level=2)
    doc.add_paragraph("Master of Science in Computer Science")
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, FastAPI, Docker, Kubernetes, PostgreSQL, Redis, Git")
    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Backend Engineer at CloudSolutions")
    
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_get_candidate_profile_default():
    """GET /resume/profile returns default profile schema before any upload."""
    response = client.get("/resume/profile")
    assert response.status_code == 200
    data = response.json()
    assert "candidate" in data
    assert "target_roles" in data
    assert data["auto_apply"] is False
    assert data["minimum_match_score"] == 75


def test_upload_unsupported_file_extension():
    """POST /resume/upload rejects unsupported extensions with 400 Bad Request."""
    invalid_file = io.BytesIO(b"Hello plain text resume")
    response = client.post(
        "/resume/upload",
        files={"file": ("resume.txt", invalid_file, "text/plain")}
    )
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]


def test_upload_empty_file():
    """POST /resume/upload rejects empty files."""
    empty_file = io.BytesIO(b"")
    response = client.post(
        "/resume/upload",
        files={"file": ("empty.pdf", empty_file, "application/pdf")}
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_upload_pdf_resume():
    """POST /resume/upload correctly parses a PDF and returns structured CandidateProfile."""
    pdf_bytes = create_sample_pdf_bytes()
    response = client.post(
        "/resume/upload",
        files={"file": ("sample_resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "sample_resume.pdf"
    assert data["extracted_text_length"] > 100
    
    profile = data["profile"]
    candidate = profile["candidate"]
    assert "Prashant" in candidate["name"]
    assert candidate["email"] == "prashant@example.com"
    assert "Python" in candidate["skills"]
    assert "FastAPI" in candidate["skills"]
    assert "Docker" in candidate["skills"]
    assert "LLM" in candidate["skills"]
    assert "AI Engineer" in profile["target_roles"]


def test_upload_docx_resume():
    """POST /resume/upload correctly parses a DOCX and returns structured CandidateProfile."""
    docx_bytes = create_sample_docx_bytes()
    response = client.post(
        "/resume/upload",
        files={"file": ("sample_resume.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "sample_resume.docx"
    assert data["extracted_text_length"] > 50

    candidate = data["profile"]["candidate"]
    assert "Alex" in candidate["name"]
    assert candidate["email"] == "alex.johnson@example.com"
    assert "Python" in candidate["skills"]
    assert "Kubernetes" in candidate["skills"]


def test_get_candidate_profile_persisted():
    """GET /resume/profile retrieves the profile saved from the last upload."""
    response = client.get("/resume/profile")
    assert response.status_code == 200
    data = response.json()
    assert data["candidate"]["email"] == "alex.johnson@example.com"


def test_patch_candidate_profile():
    """PATCH /resume/profile and PATCH /profile allow updating user preferences."""
    patch_data = {
        "target_roles": ["Principal AI Architect", "Staff ML Engineer"],
        "minimum_match_score": 85,
        "locations": ["Bangalore", "Remote", "Singapore"],
        "remote_preference": True
    }
    response = client.patch("/resume/profile", json=patch_data)
    assert response.status_code == 200
    data = response.json()
    assert data["target_roles"] == ["Principal AI Architect", "Staff ML Engineer"]
    assert data["minimum_match_score"] == 85
    assert "Singapore" in data["locations"]

    # Test alias route /profile
    alias_response = client.patch("/profile", json={"minimum_match_score": 90})
    assert alias_response.status_code == 200
    assert alias_response.json()["minimum_match_score"] == 90
